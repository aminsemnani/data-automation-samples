from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REQUIRED_COLUMNS = {
    "date",
    "order_id",
    "region",
    "salesperson",
    "product",
    "quantity",
    "unit_price",
}


def load_sales(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path)
    elif path.suffix.lower() == ".xlsx":
        frame = pd.read_excel(path)
    else:
        raise ValueError("Input must be .csv or .xlsx")

    frame.columns = [str(column).strip().lower() for column in frame.columns]
    missing = sorted(REQUIRED_COLUMNS - set(frame.columns))
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))

    frame = frame.copy()
    frame["date"] = pd.to_datetime(frame["date"], errors="raise")
    frame["quantity"] = pd.to_numeric(frame["quantity"], errors="raise")
    frame["unit_price"] = pd.to_numeric(frame["unit_price"], errors="raise")
    frame["revenue"] = frame["quantity"] * frame["unit_price"]
    frame["month"] = frame["date"].dt.to_period("M").astype(str)
    return frame.sort_values(["date", "order_id"]).reset_index(drop=True)


def summarize(frame: pd.DataFrame) -> dict:
    revenue = float(frame["revenue"].sum())
    orders = int(frame["order_id"].nunique())
    units = int(frame["quantity"].sum())
    avg_order = revenue / orders if orders else 0.0
    top_region_series = frame.groupby("region")["revenue"].sum().sort_values(ascending=False)
    top_product_series = frame.groupby("product")["revenue"].sum().sort_values(ascending=False)
    return {
        "total_revenue": round(revenue, 2),
        "orders": orders,
        "units_sold": units,
        "average_order_value": round(avg_order, 2),
        "top_region": str(top_region_series.index[0]) if not top_region_series.empty else None,
        "top_product": str(top_product_series.index[0]) if not top_product_series.empty else None,
    }


def _summary_table(frame: pd.DataFrame, group: str) -> pd.DataFrame:
    return (
        frame.groupby(group, as_index=False)
        .agg(
            orders=("order_id", "nunique"),
            units=("quantity", "sum"),
            revenue=("revenue", "sum"),
        )
        .sort_values("revenue", ascending=False)
        .reset_index(drop=True)
    )


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 42)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 10)


def _add_table(ws, table_name: str) -> None:
    if ws.max_row < 2 or ws.max_column < 1:
        return
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def build_report(input_path: Path, output_path: Path, summary_json: Path | None = None) -> dict:
    frame = load_sales(input_path)
    metrics = summarize(frame)
    monthly = _summary_table(frame, "month")
    region = _summary_table(frame, "region")
    product = _summary_table(frame, "product")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        monthly.to_excel(writer, sheet_name="Monthly", index=False)
        region.to_excel(writer, sheet_name="Region", index=False)
        product.to_excel(writer, sheet_name="Product", index=False)
        frame.drop(columns=["month"]).to_excel(writer, sheet_name="Transactions", index=False)

    wb = load_workbook(output_path)
    ws = wb.create_sheet("Executive Summary", 0)
    ws["A1"] = "Sales Management Report"
    ws["A1"].font = Font(size=18, bold=True)
    ws.merge_cells("A1:D1")
    ws["A3"] = "Metric"
    ws["B3"] = "Value"
    rows = [
        ("Total Revenue", metrics["total_revenue"]),
        ("Orders", metrics["orders"]),
        ("Units Sold", metrics["units_sold"]),
        ("Average Order Value", metrics["average_order_value"]),
        ("Top Region", metrics["top_region"]),
        ("Top Product", metrics["top_product"]),
    ]
    for row_no, (label, value) in enumerate(rows, start=4):
        ws.cell(row_no, 1, label)
        ws.cell(row_no, 2, value)
    ws["A3"].font = ws["B3"].font = Font(bold=True)
    ws["A3"].fill = ws["B3"].fill = PatternFill("solid", fgColor="D9EAF7")
    ws["B4"].number_format = "#,##0.00"
    ws["B7"].number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20

    region_ws = wb["Region"]
    chart = BarChart()
    chart.title = "Revenue by Region"
    chart.y_axis.title = "Revenue"
    chart.x_axis.title = "Region"
    data = Reference(region_ws, min_col=4, min_row=1, max_row=region_ws.max_row)
    cats = Reference(region_ws, min_col=1, min_row=2, max_row=region_ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 11
    ws.add_chart(chart, "D3")

    for sheet_name in ("Monthly", "Region", "Product", "Transactions"):
        sheet = wb[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        _autosize(sheet)

    _add_table(wb["Monthly"], "MonthlySummary")
    _add_table(wb["Region"], "RegionSummary")
    _add_table(wb["Product"], "ProductSummary")
    _add_table(wb["Transactions"], "TransactionData")

    for sheet_name in ("Monthly", "Region", "Product"):
        sheet = wb[sheet_name]
        for cell in sheet["D"][1:]:
            cell.number_format = "#,##0.00"
        sheet.conditional_formatting.add(
            f"D2:D{sheet.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="FCE8E6",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFF4CC",
                end_type="max",
                end_color="D9EAD3",
            ),
        )

    tx = wb["Transactions"]
    header_map = {cell.value: cell.column for cell in tx[1]}
    for column_name in ("unit_price", "revenue"):
        col = get_column_letter(header_map[column_name])
        for cell in tx[col][1:]:
            cell.number_format = "#,##0.00"
    date_col = get_column_letter(header_map["date"])
    for cell in tx[date_col][1:]:
        cell.number_format = "yyyy-mm-dd"
        cell.alignment = Alignment(horizontal="left")

    wb.save(output_path)

    if summary_json is not None:
        summary_json.parent.mkdir(parents=True, exist_ok=True)
        summary_json.write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build a formatted Excel management report from sales data."
    )
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--summary-json", type=Path, default=None)
    args = parser.parse_args()
    metrics = build_report(args.input, args.output, args.summary_json)
    print(json.dumps(metrics, indent=2))


if __name__ == "__main__":
    main()
