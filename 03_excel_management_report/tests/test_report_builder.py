import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from report_builder import build_report, load_sales, summarize


def test_summary_metrics_are_deterministic():
    frame = pd.DataFrame(
        {
            "date": pd.to_datetime(["2026-01-01", "2026-01-02"]),
            "order_id": ["A", "B"],
            "region": ["North", "South"],
            "salesperson": ["Alice", "Bob"],
            "product": ["P1", "P2"],
            "quantity": [2, 1],
            "unit_price": [10.0, 25.0],
            "revenue": [20.0, 25.0],
            "month": ["2026-01", "2026-01"],
        }
    )
    result = summarize(frame)
    assert result["total_revenue"] == 45.0
    assert result["orders"] == 2
    assert result["units_sold"] == 3
    assert result["average_order_value"] == 22.5
    assert result["top_region"] == "South"


def test_build_report_creates_management_workbook(tmp_path: Path):
    source = tmp_path / "sales.csv"
    source.write_text(
        "date,order_id,region,salesperson,product,quantity,unit_price\n"
        "2026-01-01,A,North,Alice,P1,2,10\n"
        "2026-01-02,B,South,Bob,P2,1,25\n",
        encoding="utf-8",
    )
    output = tmp_path / "report.xlsx"
    summary = tmp_path / "summary.json"
    metrics = build_report(source, output, summary)

    assert output.exists()
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == ["Executive Summary", "Monthly", "Region", "Product", "Transactions"]
    assert wb["Executive Summary"]["B4"].value == 45
    assert wb["Region"].tables
    assert wb["Transactions"].freeze_panes == "A2"
    assert json.loads(summary.read_text(encoding="utf-8"))["orders"] == 2
    assert metrics["top_product"] == "P2"


def test_load_sales_rejects_missing_columns(tmp_path: Path):
    source = tmp_path / "bad.csv"
    source.write_text("date,order_id\n2026-01-01,A\n", encoding="utf-8")
    try:
        load_sales(source)
    except ValueError as exc:
        assert "Missing required columns" in str(exc)
    else:
        raise AssertionError("Expected ValueError")
